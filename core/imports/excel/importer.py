from openpyxl import load_workbook

from ..base import BaseImporter
from ..exceptions import ImportFileError


class ExcelImporter(BaseImporter):

    def read(self, uploaded_file):

        try:
            workbook = load_workbook(
                uploaded_file,
                read_only=True,
                data_only=True,
            )

            worksheet = workbook.active

            rows = list(
                worksheet.iter_rows(
                    values_only=True,
                )
            )

            if not rows:
                raise ImportFileError(
                    "The Excel file is empty."
                )

            headers = [
                str(value).strip()
                if value is not None
                else ""
                for value in rows[0]
            ]

            if not any(headers):
                raise ImportFileError(
                    "The Excel file does not contain headers."
                )

            records = []

            for values in rows[1:]:
                record = dict(
                    zip(headers, values)
                )

                if any(
                    value not in (None, "")
                    for value in record.values()
                ):
                    records.append(record)

            return records

        except Exception as exc:
            if isinstance(exc, ImportFileError):
                raise

            raise ImportFileError(
                "Unable to read the Excel file."
            ) from exc